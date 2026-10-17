"""Builds outputs/P2P_Risk_Control_Matrix.xlsx and the process-map figure."""
import sys
from pathlib import Path
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

sys.path.insert(0, str(Path(__file__).parent))
from chartstyle import apply, C, save
sys.path.insert(0, str(Path(__file__).parent))
import rcm_def as D

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"; OUT.mkdir(exist_ok=True)
PROC = ROOT / "data" / "processed"

# ------------------------------------------------------------------ process map
apply()
steps = [("Purchase\nrequisition", "R1"), ("Purchase\norder", "R2, R5"), ("Vendor\nset-up", "R3"), ("Goods receipt /\nservice entry", "R4"),
         ("Invoice\nreceipt", "R6, R7"), ("Payment\nclearing", "R8, R9, R10")]
fig, ax = plt.subplots(figsize=(10, 2.6)); ax.set_xlim(0, 60); ax.set_ylim(0, 16); ax.axis("off")
for k, (name, risks) in enumerate(steps):
    x = 1 + k * 9.7
    ax.add_patch(FancyBboxPatch((x, 7), 8, 6, boxstyle="round,pad=0.2", fc="#E8EEF6", ec="#1F3A5F", lw=1.2))
    ax.text(x + 4, 10, name, ha="center", va="center", fontsize=8.5, color="#1F3A5F", fontweight="bold")
    ax.text(x + 4, 4.5, risks, ha="center", va="center", fontsize=8, color="#B42318", fontweight="bold")
    if k < len(steps) - 1:
        ax.annotate("", xy=(x + 9.5, 10), xytext=(x + 8.3, 10), arrowprops=dict(arrowstyle="-|>", color="#1F3A5F", lw=1.2))
ax.text(0.5, 1.6, "Red labels are the risk IDs in the matrix. Vendor set-up sits outside the log, so R3 is documented but not tested.", fontsize=7.5, color="#555555")
save(fig, str(OUT / "process_map.png")); plt.close(fig)

# ------------------------------------------------------------------ workbook
wb = Workbook()
HDR = PatternFill("solid", fgColor="1F3A5F"); WH = Font(bold=True, color="FFFFFF"); thin = Side(style="thin", color="C9D0DA")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

ws = wb.active; ws.title = "RCM"
cols = ["Risk ID", "Process step", "Risk", "Category", "Likelihood (1-5)", "Impact (1-5)", "Inherent score", "Inherent rating",
        "Control ID", "Control", "Type", "Nature", "Frequency", "Control owner", "Assertion", "Test ID", "Test procedure",
        "Population", "Exceptions", "Exception rate", "Control conclusion", "Residual risk", "Value at risk (item value)"]
ws.append(cols)
for c in ws[1]:
    c.fill, c.font, c.alignment, c.border = HDR, WH, Alignment(wrap_text=True, vertical="center"), BOX
valmap = {"T1a": "T1a", "T1b": "T1b", "T2": "T2", "T4": "T4", "T5": "T5"}
for i, r in enumerate(D.ROWS, start=2):
    cnt = D.counts(r); tested = cnt is not None
    pop, exc = cnt if tested else ("", "")
    val = float(D.RES.loc[r[13], "value"]) if r[13] in valmap else ""
    ws.append([r[0], r[1], r[2], r[3], r[4], r[5], f"=E{i}*F{i}", f'=IF(G{i}>=15,"High",IF(G{i}>=8,"Medium","Low"))',
               r[6], r[7], r[8], r[9], r[10], r[11], r[12], r[13] or "n/a", r[14],
               pop, exc, f'=IF(R{i}="","",S{i}/R{i})',
               f'=IF(R{i}="","Not tested",IF(T{i}<=0.01,"Effective",IF(T{i}<=0.05,"Needs improvement","Ineffective")))',
               f'=IF(U{i}="Effective",IF(H{i}="High","Medium","Low"),H{i})', val])
    ws.cell(i, 20).number_format = "0.0%"; ws.cell(i, 23).number_format = "#,##0"
    ws.cell(i, 18).number_format = "#,##0"; ws.cell(i, 19).number_format = "#,##0"
# R5 value (T3) and R4 value come from the SoD exception table
ex3 = pd.read_parquet(PROC / "ex_T3.parquet")
v4 = float(ex3[ex3.detail.str.contains("goods receipt")].item_value.sum())
v5 = float(ex3[ex3.detail.str.contains("invoice|payment")].item_value.sum())
for i, r in enumerate(D.ROWS, start=2):
    if r[0] == "R4": ws.cell(i, 23).value = v4
    if r[0] == "R5": ws.cell(i, 23).value = v5
widths = [8, 14, 36, 11, 11, 9, 9, 11, 8, 40, 14, 14, 13, 18, 15, 8, 44, 11, 11, 10, 17, 11, 16]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(j)].width = w
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for c in row:
        c.alignment, c.border = WRAP, BOX
ws.row_dimensions[1].height = 34; ws.freeze_panes = "D2"
rng = f"U2:U{ws.max_row}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Effective"'], fill=PatternFill("solid", bgColor="D1FADF")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Needs improvement"'], fill=PatternFill("solid", bgColor="FEF0C7")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Ineffective"'], fill=PatternFill("solid", bgColor="FEE4E2")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Not tested"'], fill=PatternFill("solid", bgColor="E4E7EC")))

# rating key
rk = wb.create_sheet("Rating key")
for r in [("Measure", "Rule"),
          ("Inherent score", "Likelihood x impact, each 1 to 5"),
          ("Inherent rating", "High 15 to 25 | Medium 8 to 14 | Low 1 to 7"),
          ("Control conclusion", "Exception rate up to 1% Effective | above 1% to 5% Needs improvement | above 5% Ineffective | no test: Not tested"),
          ("Residual risk", "Effective control lowers inherent rating one level (Low stays Low); any other result leaves it unchanged"),
          ("Likelihood and impact", "Reviewer judgement from the process and the nature of the exposure; not derived from the data")]:
    rk.append(r)
for c in rk[1]: c.fill, c.font = HDR, WH
rk.column_dimensions["A"].width = 24; rk.column_dimensions["B"].width = 110

# test results
tr = wb.create_sheet("Test results")
tr.append(["Test", "Description", "SQL file", "Population", "Exceptions", "Rate", "Item value of exceptions"])
desc = {"T1a": ("Invoice cleared with no goods receipt", "sql/T1a_paid_without_goods_receipt.sql"),
        "T1b": ("GR-based invoice receipt before goods receipt", "sql/T1b_invoice_before_receipt_on_gr_based_items.sql"),
        "T2": ("Duplicate payment leads", "sql/T2_duplicate_payments.sql"),
        "T3": ("Segregation of duties (any conflict)", "sql/T3_segregation_of_duties.sql"),
        "T4": ("Payment before goods receipt", "sql/T4_paid_before_goods_receipt.sql"),
        "T5": ("Split purchase orders", "sql/T5_split_purchase_orders.sql")}
for t, r in D.RES.iterrows():
    tr.append([t, desc[t][0], desc[t][1], int(r.population), int(r.exceptions), f"=E{tr.max_row+1}/D{tr.max_row+1}", float(r.value)])
for c in tr[1]: c.fill, c.font = HDR, WH
for i in range(2, tr.max_row + 1):
    tr.cell(i, 6).number_format = "0.0%"
    for j in (4, 5, 7): tr.cell(i, j).number_format = "#,##0"
for j, w in enumerate([8, 44, 56, 12, 12, 9, 22], 1): tr.column_dimensions[get_column_letter(j)].width = w

# exception samples
def sample(name, t, n=300):
    d = pd.read_parquet(PROC / f"ex_{t}.parquet").sort_values("item_value", ascending=False).head(n)
    for col in d.columns:
        if str(d[col].dtype).startswith("datetime"): d[col] = d[col].dt.strftime("%Y-%m-%d %H:%M")
    s = wb.create_sheet(name); s.append(list(d.columns))
    for c in s[1]: c.fill, c.font = HDR, WH
    for row in d.itertuples(index=False): s.append([None if pd.isna(v) else v for v in row])
    for j in range(1, len(d.columns) + 1): s.column_dimensions[get_column_letter(j)].width = 22
    s.column_dimensions[get_column_letter(list(d.columns).index("detail") + 1)].width = 70
    s.freeze_panes = "A2"
for nm, t in [("Ex T1a", "T1a"), ("Ex T1b", "T1b"), ("Ex T2", "T2"), ("Ex T3", "T3"), ("Ex T4", "T4"), ("Ex T5", "T5")]:
    sample(nm, t)

# scope
sc = wb.create_sheet("Scope and limits")
for r in [("Item", "Detail"),
          ("Data", f"BPI Challenge 2019 purchase-order event log (4TU.ResearchData, CC BY 4.0): {D.ST['cases']:,} purchase items, {D.ST['events']:,} events. Real data from a multinational; company, vendor and user names and monetary values are anonymised."),
          ("Value used", "The log holds one cumulative item value per event, not separate goods receipt and invoice amounts. Three-way match is tested on event sequence; value at risk is the item value."),
          ("Not in the log", "Vendor master changes, approval matrix, invoice numbers, payment amounts per invoice. R3 is documented but not tested; the approval limit in T5 is an assumption."),
          ("Exception samples", "Each Ex sheet shows the 300 highest-value exceptions; the full tables are rebuilt by src/run_tests.py."),
          ("Rating", "See the Rating key sheet.")]:
    sc.append(r)
for c in sc[1]: c.fill, c.font = HDR, WH
sc.column_dimensions["A"].width = 20; sc.column_dimensions["B"].width = 120
for row in sc.iter_rows(min_row=2):
    for c in row: c.alignment = WRAP
wb.properties.creator = "Shaswat Sharma"
wb.save(OUT / "P2P_Risk_Control_Matrix.xlsx")
print("rcm built")
